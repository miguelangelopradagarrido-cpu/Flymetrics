import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Asigna color de fondo a una celda de tabla en Word."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Define márgenes internos de una celda en dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = parse_xml(f'<w:{m} {nsdecls("w")} w:w="{val}" w:type="dxa"/>')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_bitacora_docx(data: dict) -> io.BytesIO:
    """
    Genera un documento Word (.docx) oficial, elegante y estructurado
    de la Bitácora de Vuelo RPAS conforme al RAC-100 y manual de marca Flymetrics.
    """
    doc = Document()
    
    # Configurar márgenes de página (estrechos para formato técnico A4)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)
        s.page_width = Inches(8.27)  # A4
        s.page_height = Inches(11.69)

    navy_color = RGBColor(0, 48, 73)      # #003049
    blue_color = RGBColor(28, 130, 173)   # #1C82AD
    gold_color = RGBColor(200, 169, 106)  # #C8A96A
    gray_muted = RGBColor(100, 116, 139)  # #64748B
    dark_text = RGBColor(51, 65, 85)      # #334155

    # 1. ENCABEZADO CORPORATIVO
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(4.8)
    header_table.columns[1].width = Inches(2.2)

    # Celda izquierda: Nombre de marca
    c_left = header_table.cell(0, 0)
    set_cell_background(c_left, "003049")
    set_cell_margins(c_left, top=140, bottom=140, left=180, right=140)
    p_brand = c_left.paragraphs[0]
    p_brand.paragraph_format.space_before = Pt(0)
    p_brand.paragraph_format.space_after = Pt(2)
    r_brand = p_brand.add_run("FLYMETRICS S.A.S.")
    r_brand.font.bold = True
    r_brand.font.size = Pt(16)
    r_brand.font.color.rgb = RGBColor(255, 255, 255)
    
    p_sub = c_left.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(0)
    r_sub = p_sub.add_run("GESTIÓN INTELIGENTE DE DRONES AGRÍCOLAS & RPAS | COLOMBIA")
    r_sub.font.size = Pt(8)
    r_sub.font.color.rgb = RGBColor(200, 169, 106)
    r_sub.font.bold = True

    # Celda derecha: Radicado y certificación
    c_right = header_table.cell(0, 1)
    set_cell_background(c_right, "0a4a6e")
    set_cell_margins(c_right, top=140, bottom=140, left=140, right=180)
    p_rad = c_right.paragraphs[0]
    p_rad.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_rad.paragraph_format.space_before = Pt(0)
    p_rad.paragraph_format.space_after = Pt(2)
    r_rad_label = p_rad.add_run("BITÁCORA OFICIAL\n")
    r_rad_label.font.size = Pt(8)
    r_rad_label.font.color.rgb = RGBColor(255, 255, 255)
    r_rad_label.font.bold = True
    
    rad_num = data.get("rad", f"RAD-2026-{str(data.get('id_turno', '001')).zfill(3)}")
    r_rad_val = p_rad.add_run(f"#{rad_num}")
    r_rad_val.font.size = Pt(14)
    r_rad_val.font.bold = True
    r_rad_val.font.color.rgb = RGBColor(200, 169, 106)

    # Espacio
    p_sep = doc.add_paragraph()
    p_sep.paragraph_format.space_before = Pt(6)
    p_sep.paragraph_format.space_after = Pt(4)

    # 2. TÍTULO DEL DOCUMENTO
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_doc_title = p_title.add_run("BITÁCORA DE VUELO Y APLICACIÓN RPAS (RAC-100)")
    r_doc_title.font.bold = True
    r_doc_title.font.size = Pt(12)
    r_doc_title.font.color.rgb = navy_color

    p_norma = doc.add_paragraph()
    p_norma.paragraph_format.space_before = Pt(0)
    p_norma.paragraph_format.space_after = Pt(8)
    r_norma = p_norma.add_run("Conforme a la Reglamentación Aeronáutica de Colombia RAC 91 / RAC 100 — Unidad Administrativa Especial de Aeronáutica Civil")
    r_norma.font.size = Pt(8)
    r_norma.font.color.rgb = gray_muted

    def add_section_header(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f"■ {title_text.upper()}")
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = blue_color

    # 3. SECCIÓN: INFORMACIÓN DEL CLIENTE Y PREDIO
    add_section_header("1. Información del Cliente y del Predio")
    t_cli = doc.add_table(rows=3, cols=4)
    t_cli.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_cli.autofit = False
    
    col_widths = [Inches(1.8), Inches(2.2), Inches(1.4), Inches(1.6)]
    for row in t_cli.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    fields_cli = [
        ("Cliente / Titular:", data.get("cliente_nombre", "—"), "Cédula / NIT:", data.get("cliente_cedula", "—")),
        ("Teléfono Contacto:", data.get("cliente_tel", "—"), "Correo:", data.get("cliente_email", "—")),
        ("Nombre del Predio:", data.get("finca_nombre", "—"), "Ubicación:", f"{data.get('finca_municipio', 'Yopal')} ({data.get('departamento', 'Casanare')})")
    ]

    for r_idx, (l1, v1, l2, v2) in enumerate(fields_cli):
        c0, c1, c2, c3 = t_cli.cell(r_idx, 0), t_cli.cell(r_idx, 1), t_cli.cell(r_idx, 2), t_cli.cell(r_idx, 3)
        for c in [c0, c2]: set_cell_background(c, "F1F5F9")
        for c in [c0, c1, c2, c3]: set_cell_margins(c, 60, 60, 80, 80)
        
        p0 = c0.paragraphs[0]; p0.add_run(l1).font.bold = True; p0.runs[0].font.size = Pt(8); p0.runs[0].font.color.rgb = navy_color
        p1 = c1.paragraphs[0]; p1.add_run(str(v1)).font.size = Pt(8.5); p1.runs[0].font.color.rgb = dark_text
        p2 = c2.paragraphs[0]; p2.add_run(l2).font.bold = True; p2.runs[0].font.size = Pt(8); p2.runs[0].font.color.rgb = navy_color
        p3 = c3.paragraphs[0]; p3.add_run(str(v2)).font.size = Pt(8.5); p3.runs[0].font.color.rgb = dark_text

    # 4. SECCIÓN: DATOS DE OPERACIÓN Y AERONAVE
    add_section_header("2. Datos Técnicos de Vuelo y Aeronave RPAS")
    t_ops = doc.add_table(rows=3, cols=4)
    t_ops.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_ops.autofit = False
    for row in t_ops.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    fields_ops = [
        ("Fecha de Operación:", str(data.get("fecha_vuelo", "—")), "Horario de Vuelo:", f"{data.get('hora_ini', '07:00')} - {data.get('hora_fin', '09:00')}"),
        ("Aeronave RPAS / Dron:", data.get("drone_modelo", "DJI Agras T40"), "Matrícula / Serie:", data.get("drone_serie", "T40-CO-8942")),
        ("Piloto en Mando:", data.get("piloto_nombre", "Piloto Certificado Aerocivil"), "C.C. / Licencia RPAS:", f"{data.get('piloto_cedula', '1.121.890.345')} ({data.get('piloto_licencia', 'RPAS-COL-8942')})")
    ]

    for r_idx, (l1, v1, l2, v2) in enumerate(fields_ops):
        c0, c1, c2, c3 = t_ops.cell(r_idx, 0), t_ops.cell(r_idx, 1), t_ops.cell(r_idx, 2), t_ops.cell(r_idx, 3)
        for c in [c0, c2]: set_cell_background(c, "F1F5F9")
        for c in [c0, c1, c2, c3]: set_cell_margins(c, 60, 60, 80, 80)
        p0 = c0.paragraphs[0]; p0.add_run(l1).font.bold = True; p0.runs[0].font.size = Pt(8); p0.runs[0].font.color.rgb = navy_color
        p1 = c1.paragraphs[0]; p1.add_run(str(v1)).font.size = Pt(8.5); p1.runs[0].font.color.rgb = dark_text
        p2 = c2.paragraphs[0]; p2.add_run(l2).font.bold = True; p2.runs[0].font.size = Pt(8); p2.runs[0].font.color.rgb = navy_color
        p3 = c3.paragraphs[0]; p3.add_run(str(v2)).font.size = Pt(8.5); p3.runs[0].font.color.rgb = dark_text

    # 5. SECCIÓN: MÉTRICAS AGRONÓMICAS Y APLICACIÓN ULV
    add_section_header("3. Métricas Agronómicas y Balance de Mezcla")
    t_met = doc.add_table(rows=2, cols=4)
    t_met.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_met.autofit = False
    for row in t_met.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    fields_met = [
        ("Área Total Cubierta:", f"{data.get('finca_area', data.get('area_ha', '10.0'))} ha", "Insumo / Agroquímico:", f"{data.get('insumos_litros', '15.0')} Litros"),
        ("Agua / Portador:", f"{data.get('agua_litros', '100.0')} Litros", "Baterías Empleadas:", f"{data.get('baterias', data.get('baterias_utilizadas', '4'))} Ciclos Smart")
    ]

    for r_idx, (l1, v1, l2, v2) in enumerate(fields_met):
        c0, c1, c2, c3 = t_met.cell(r_idx, 0), t_met.cell(r_idx, 1), t_met.cell(r_idx, 2), t_met.cell(r_idx, 3)
        for c in [c0, c2]: set_cell_background(c, "EBF5FB")
        for c in [c0, c1, c2, c3]: set_cell_margins(c, 60, 60, 80, 80)
        p0 = c0.paragraphs[0]; p0.add_run(l1).font.bold = True; p0.runs[0].font.size = Pt(8); p0.runs[0].font.color.rgb = navy_color
        p1 = c1.paragraphs[0]; p1.add_run(str(v1)).font.size = Pt(9); p1.runs[0].font.bold = True; p1.runs[0].font.color.rgb = blue_color
        p2 = c2.paragraphs[0]; p2.add_run(l2).font.bold = True; p2.runs[0].font.size = Pt(8); p2.runs[0].font.color.rgb = navy_color
        p3 = c3.paragraphs[0]; p3.add_run(str(v2)).font.size = Pt(9); p3.runs[0].font.bold = True; p3.runs[0].font.color.rgb = blue_color

    # 6. SECCIÓN: OBSERVACIONES Y RECOMENDACIONES
    add_section_header("4. Novedades Operativas y Recomendaciones de Campo")
    t_obs = doc.add_table(rows=2, cols=1)
    t_obs.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_obs.autofit = False
    t_obs.columns[0].width = Inches(7.0)

    c_obs = t_obs.cell(0, 0)
    set_cell_background(c_obs, "F8FAFC")
    set_cell_margins(c_obs, 80, 80, 100, 100)
    p_obs = c_obs.paragraphs[0]
    p_obs.add_run("Observaciones de Vuelo y Condiciones Climáticas:\n").font.bold = True
    p_obs.runs[0].font.size = Pt(8); p_obs.runs[0].font.color.rgb = navy_color
    r_obs_val = p_obs.add_run(data.get("observaciones_campo", "Operación ejecutada con viento inferior a 10 km/h, cielo despejado y sin deriva hacia fuentes hídricas."))
    r_obs_val.font.size = Pt(8.5); r_obs_val.font.color.rgb = dark_text

    c_rec = t_obs.cell(1, 0)
    set_cell_background(c_rec, "F0FDF4")
    set_cell_margins(c_rec, 80, 80, 100, 100)
    p_rec = c_rec.paragraphs[0]
    p_rec.add_run("Recomendaciones Agronómicas Flymetrics:\n").font.bold = True
    p_rec.runs[0].font.size = Pt(8); p_rec.runs[0].font.color.rgb = RGBColor(22, 101, 52)
    r_rec_val = p_rec.add_run(data.get("recomendaciones_agronomicas", "Monitorear absorción foliar a las 48 horas. Realizar evaluación visual de eficacia en 10 días."))
    r_rec_val.font.size = Pt(8.5); r_rec_val.font.color.rgb = RGBColor(21, 128, 61)

    # 7. ENLACE A SHAREPOINT SI APLICA
    sharepoint_url = data.get("url_entregable_drive") or data.get("link_sharepoint")
    if sharepoint_url:
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(6)
        p_sp.paragraph_format.space_after = Pt(2)
        r_sp_l = p_sp.add_run("🔗 Enlace Oficial al Informe y Ortofoto en SharePoint: ")
        r_sp_l.font.bold = True
        r_sp_l.font.size = Pt(8)
        r_sp_l.font.color.rgb = blue_color
        r_sp_u = p_sp.add_run(str(sharepoint_url))
        r_sp_u.font.size = Pt(8)
        r_sp_u.font.underline = True

    # 8. SECCIÓN DE FIRMAS Y CONFORMIDAD
    add_section_header("5. Firmas de Conformidad y Aprobación")
    t_firmas = doc.add_table(rows=1, cols=3)
    t_firmas.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_firmas.autofit = False
    f_width = Inches(2.33)
    for c in t_firmas.rows[0].cells:
        c.width = f_width
        set_cell_margins(c, 100, 100, 80, 80)
        set_cell_background(c, "F8FAFC")

    # Firma 1: Piloto
    c_f1 = t_firmas.cell(0, 0)
    p_f1 = c_f1.paragraphs[0]
    p_f1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_f1.add_run("_____________________________\n").font.size = Pt(8)
    r1 = p_f1.add_run(f"{data.get('piloto_nombre', 'Piloto en Mando')}\n")
    r1.font.bold = True; r1.font.size = Pt(8.5); r1.font.color.rgb = navy_color
    p_f1.add_run("Piloto Certificado Aerocivil\nC.C. " + str(data.get('piloto_cedula', '—'))).font.size = Pt(7.5)

    # Firma 2: Supervisor Flymetrics
    c_f2 = t_firmas.cell(0, 1)
    p_f2 = c_f2.paragraphs[0]
    p_f2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_f2.add_run("_____________________________\n").font.size = Pt(8)
    r2 = p_f2.add_run("Supervisor Técnico Flymetrics\n")
    r2.font.bold = True; r2.font.size = Pt(8.5); r2.font.color.rgb = navy_color
    p_f2.add_run("Control de Operaciones RPAS\nC.C. 1.053.489.120").font.size = Pt(7.5)

    # Firma 3: Cliente / Productor
    c_f3 = t_firmas.cell(0, 2)
    p_f3 = c_f3.paragraphs[0]
    p_f3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_f3.add_run("_____________________________\n").font.size = Pt(8)
    r3 = p_f3.add_run(f"{data.get('cliente_nombre', 'Propietario / Cliente')}\n")
    r3.font.bold = True; r3.font.size = Pt(8.5); r3.font.color.rgb = navy_color
    p_f3.add_run("Propietario / Representante\nNIT / C.C. " + str(data.get('cliente_cedula', '—'))).font.size = Pt(7.5)

    # 9. PIE DE PÁGINA Y AVISO LEGAL
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(12)
    p_foot.paragraph_format.space_after = Pt(0)
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = p_foot.add_run(
        "Flymetrics S.A.S. — Yopal, Casanare, Colombia | Tel: +57 305 406 1764 | Email: soporte@flymetrics.co\n"
        "Documento emitido electrónicamente bajo estándares de seguridad operacional RPAS y certificación RAC-100."
    )
    r_foot.font.size = Pt(7)
    r_foot.font.color.rgb = gray_muted

    # Guardar en buffer en memoria
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
